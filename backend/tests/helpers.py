"""Test helper utilities for generating valid and edge-case document fixtures."""

import os
import docx
import pypdf
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject


def create_test_txt(directory: str, filename: str = "test_policy.txt", content: str = "Paragraph 1.\n\nParagraph 2.") -> str:
    """Generates a sample TXT document."""
    os.makedirs(directory, exist_ok=True)
    file_path = os.path.join(directory, filename)
    with open(file_path, "w", encoding="utf-8") as f:
        f.write(content)
    return file_path


def create_test_docx(
    directory: str,
    filename: str = "test_policy.docx",
    sections: list[tuple[str, list[str]]] | None = None,
) -> str:
    """Generates a sample DOCX document with headings and paragraphs."""
    os.makedirs(directory, exist_ok=True)
    file_path = os.path.join(directory, filename)
    doc = docx.Document()

    if sections is None:
        sections = [
            ("Paid Time Off Policy", ["Employees receive 20 days of paid leave per year.", "Unused leave rolls over up to 5 days."]),
            ("Bereavement Policy", ["Up to 3 days of paid leave are granted for immediate family members."]),
        ]

    for heading, paragraphs in sections:
        doc.add_heading(heading, level=1)
        for p in paragraphs:
            doc.add_paragraph(p)

    doc.save(file_path)
    return file_path


def create_test_pdf(
    directory: str,
    filename: str = "test_policy.pdf",
    blank_pages: int = 0,
    text_pages: list[str] | None = None,
) -> str:
    """Generates a valid PDF document with optional text or blank pages."""
    os.makedirs(directory, exist_ok=True)
    file_path = os.path.join(directory, filename)
    writer = pypdf.PdfWriter()

    # Reusable font reference
    font_dict = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    font_ref = writer._add_object(font_dict)

    # 1. Blank pages (simulating scanned / empty pages)
    for _ in range(blank_pages):
        writer.add_blank_page(width=300, height=300)

    # 2. Text pages with content streams
    if text_pages:
        for text in text_pages:
            page = writer.add_blank_page(width=300, height=300)
            if "/Resources" not in page:
                page[NameObject("/Resources")] = DictionaryObject()
            if "/Font" not in page["/Resources"]:
                page["/Resources"][NameObject("/Font")] = DictionaryObject()
            page["/Resources"][NameObject("/Font")][NameObject("/F1")] = font_ref

            stream_obj = DecodedStreamObject()
            stream_obj.set_data(f"BT /F1 12 Tf 50 250 Td ({text}) Tj ET".encode("latin-1"))
            page[NameObject("/Contents")] = writer._add_object(stream_obj)

    # If neither specified, default to 1 text page
    if blank_pages == 0 and not text_pages:
        page = writer.add_blank_page(width=300, height=300)
        if "/Resources" not in page:
            page[NameObject("/Resources")] = DictionaryObject()
        if "/Font" not in page["/Resources"]:
            page["/Resources"][NameObject("/Font")] = DictionaryObject()
        page["/Resources"][NameObject("/Font")][NameObject("/F1")] = font_ref

        stream_obj = DecodedStreamObject()
        stream_obj.set_data(b"BT /F1 12 Tf 50 250 Td (Default Policy Content) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream_obj)

    with open(file_path, "wb") as f:
        writer.write(f)

    return file_path
